"""Build the static GitHub Pages site (docs/) from scraped prices and the local booking file.

Usage:
    python src/build_site.py
    python src/build_site.py --flights data/archive/flights_2026.csv --out <dir>   # preview an old season
"""
import argparse
import glob
import json
import os
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook

from functions import get_dates_between, load_config

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
TEMPLATE_PATH = os.path.join(ROOT, 'site', 'index.html')
DATA_PLACEHOLDER = '/*__FLIGHT_DATA__*/null'
# Days-before-departure buckets for the "When to buy" tab, earliest first: (from, to), to=None means "or more"
LEAD_BUCKETS = [(240, None), (180, 239), (150, 179), (120, 149), (90, 119), (75, 89), (60, 74),
                (45, 59), (30, 44), (21, 29), (14, 20), (7, 13), (3, 6), (1, 2), (0, 0)]


def to_iso(value):
    """Excel cells may hold datetimes, dates or text; return 'YYYY-MM-DD' or None."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.strip():
        for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d. %m. %Y', '%d/%m/%Y'):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def load_stays(booking_path):
    """Read stays (arrival/departure) from the local booking workbook.

    Names and notes are deliberately not returned - only dates are published.
    """
    if not os.path.exists(booking_path):
        print(f"Booking file not found ({booking_path}); house shown as free.")
        return []

    ws = load_workbook(booking_path, data_only=True, read_only=True)['Stays']
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h else '' for h in next(rows)]
    arrival_col, departure_col = header.index('arrival'), header.index('departure')

    stays = []
    for row_number, row in enumerate(rows, start=2):
        arrival, departure = to_iso(row[arrival_col]), to_iso(row[departure_col])
        if not arrival and not departure:
            continue
        if not arrival or not departure or departure <= arrival:
            print(f"Skipping booking row {row_number}: check arrival/departure dates.")
            continue
        stays.append({'arrival': arrival, 'departure': departure})
    return sorted(stays, key=lambda s: s['arrival'])


def load_prices(flights_path, flight_dates):
    """Return offer dates and per-flight price history for both directions."""
    records = pd.read_csv(flights_path)
    offers = sorted(records['offer_date'].unique().tolist())

    def history(direction):
        rows = records[records['direction'] == direction].set_index('offer_date').reindex(offers)
        result = {}
        for flight in flight_dates:
            column = rows[flight] if flight in rows.columns else pd.Series(index=offers, dtype=float)
            result[flight] = [None if pd.isna(p) else int(round(float(p))) for p in column]
        return result

    return offers, history('From Prague'), history('To Prague')


def load_history(archive_dir, before_year):
    """Median fare by days before departure in past seasons, per direction and flight month.

    Uses archived seasons before `before_year` that were tracked from at least 180 days out
    (flights_2024.csv only covers the last two months, which would skew the short-notice rows).
    """
    frames = []
    for path in sorted(glob.glob(os.path.join(archive_dir, 'flights_*.csv'))):
        records = pd.read_csv(path)
        prices = records.melt(id_vars=['offer_date', 'direction'], var_name='flight', value_name='price').dropna()
        prices['flight'] = pd.to_datetime(prices['flight'])
        prices['lead'] = (prices['flight'] - pd.to_datetime(prices['offer_date'])).dt.days
        prices = prices[prices['lead'] >= 0]
        if prices.empty or prices['lead'].max() < 180 or prices['flight'].dt.year.min() >= before_year:
            continue
        prices['season'] = prices['flight'].dt.year
        frames.append(prices)
    if not frames:
        return None
    prices = pd.concat(frames)

    def cell(rows):
        if rows.empty:
            return None
        return {'median': int(round(rows['price'].median())), 'flights': int(rows['flight'].nunique()),
                'seasons': {str(y): int(round(p)) for y, p in rows.groupby('season')['price'].median().items()}}

    table = {}
    for key, direction in (('out', 'From Prague'), ('back', 'To Prague')):
        legs = prices[prices['direction'] == direction]
        table[key] = {}
        for month, rows in legs.groupby(legs['flight'].dt.month):
            table[key][str(month)] = [
                cell(rows[rows['lead'].between(lo, 10_000 if hi is None else hi)])
                for lo, hi in LEAD_BUCKETS]
    weekdays = sorted(set(prices['flight']), key=lambda d: d.weekday())
    return {
        'seasons': sorted(int(y) for y in prices['season'].unique()),
        'days': list(dict.fromkeys(d.strftime('%A') for d in weekdays)),
        'buckets': LEAD_BUCKETS,
        'out': table['out'],
        'back': table['back'],
    }


def build(flights_path, config, booking_path, out_dir):
    flight_dates = get_dates_between(
        start_date=config['flight_dates_from'],
        end_date=config['flight_dates_to'],
        days_of_week=config['flight_days'])
    offers, outbound, inbound = load_prices(flights_path, flight_dates)

    data = {
        'route': {'from': 'PRG', 'to': 'JSI', 'fromName': 'Prague', 'toName': 'Skiathos'},
        'season': {
            'from': config['flight_dates_from'],
            'to': config['flight_dates_to'],
            'days': config['flight_days'],
        },
        'flights': flight_dates,
        'offers': offers,
        'out': outbound,
        'back': inbound,
        'stays': load_stays(booking_path),
        'history': load_history(os.path.join(ROOT, 'data', 'archive'), int(config['flight_dates_from'][:4])),
    }

    os.makedirs(out_dir, exist_ok=True)
    payload = json.dumps(data, separators=(',', ':'))
    with open(TEMPLATE_PATH, encoding='utf-8') as f:
        html = f.read().replace(DATA_PLACEHOLDER, payload.replace('</', '<\\/'))
    with open(os.path.join(out_dir, 'index.html'), 'w', encoding='utf-8') as f:
        f.write(html)
    with open(os.path.join(out_dir, 'data.json'), 'w', encoding='utf-8') as f:
        f.write(payload)
    # Serve files as-is, without Jekyll processing
    open(os.path.join(out_dir, '.nojekyll'), 'w').close()

    print(f"Site built in {out_dir}: {len(flight_dates)} flights, {len(offers)} price checks, "
          f"{len(data['stays'])} stays.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--flights', default=os.path.join(ROOT, 'data', 'flights.csv'))
    parser.add_argument('--config', default=os.path.join(ROOT, 'config', 'config.json'))
    parser.add_argument('--booking', default=os.path.join(ROOT, 'booking', 'booking.xlsx'))
    parser.add_argument('--out', default=os.path.join(ROOT, 'docs'))
    parser.add_argument('--season-from', help='Override flight_dates_from (for previewing old data)')
    parser.add_argument('--season-to', help='Override flight_dates_to')
    parser.add_argument('--days', nargs='+', help='Override flight_days, e.g. Monday Thursday')
    args = parser.parse_args()

    config = load_config(args.config)
    if args.season_from: config['flight_dates_from'] = args.season_from
    if args.season_to: config['flight_dates_to'] = args.season_to
    if args.days: config['flight_days'] = args.days

    build(args.flights, config, args.booking, args.out)


if __name__ == '__main__':
    main()
